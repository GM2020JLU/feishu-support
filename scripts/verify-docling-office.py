"""Real optional Docling conversion canary; synthetic documents, no live data.

Run with the separately locked Office environment's Python interpreter.
"""

import contextlib
import io
import json
import sys
import sysconfig
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from k3_support.docling_cli import main as attachment_main
from k3_support.docling_convert import convert_office
from k3_support.docling_review import preview
from k3_support.docling_draft import build_draft
from k3_support.knowledge_authoring import save_attachment, detail, export_markdown
from k3_support.db import connect, migrate


def verify_saved_draft(evidence):
    fields = {'evidence': evidence, 'title': 'Synthetic attachment review',
        'question': 'Where is the synthetic example?', 'answer': 'Synthetic test only; do not execute.',
        'references': [evidence['reading_order']['body'][0]],
        'risk_class': 'read_only', 'rollback': ''}
    draft = build_draft(**fields)
    with tempfile.TemporaryDirectory(prefix='codex-docling-draft-') as directory:
        database = Path(directory) / 'synthetic.db'
        conn = connect(database)
        try:
            migrate(conn)
            receipt = save_attachment(conn, fields=fields,
                expected_digest=draft['revision_digest'], actor_id='synthetic-reviewer')
        finally:
            conn.close()
        conn = connect(database)
        try:
            before = conn.serialize()
            saved = detail(conn, candidate_id=receipt['candidate_id'])
            target = Path(directory) / 'draft.md'
            export_markdown(conn, candidate_id=receipt['candidate_id'], output=target)
            if (saved['markdown'] != draft['markdown'] or target.read_text() != draft['markdown']
                    or saved['status'] != 'captured' or saved['reviewed']
                    or saved['material']['automatic_reply_eligible']
                    or target.stat().st_mode & 0o777 != 0o600
                    or conn.serialize() != before):
                raise ValueError('saved draft/export changed content, authority or database')
            if conn.execute('SELECT count(*) FROM knowledge_entries').fetchone()[0]:
                raise ValueError('draft unexpectedly entered runtime knowledge')
        finally:
            conn.close()


def fixtures():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    doc = Document()
    doc.add_heading("Synthetic support procedure", 0)
    doc.add_paragraph("SYNTHETIC-FAN-CANARY")
    doc.add_heading("合成命令示例（不要执行）", level=1)
    code = doc.add_paragraph()
    code.add_run('printenv bootcmd\nprintenv bootargs\necho "${fixture_value}"').font.name = 'Courier New'
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "PWM-CANARY"
    table.cell(1, 1).text = "128"
    data = io.BytesIO()
    doc.save(data)
    yield "docx", data.getvalue(), "PWM-CANARY"

    book = Workbook()
    book.active.append(["Parameter", "Value"])
    book.active.append(["PWM-CANARY", 128])
    data = io.BytesIO()
    book.save(data)
    yield "xlsx", data.getvalue(), "PWM-CANARY"

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "SYNTHETIC-FAN-CANARY"
    slide.placeholders[1].text = "PWM-CANARY 128"
    data = io.BytesIO()
    slides.save(data)
    yield "pptx", data.getvalue(), "PWM-CANARY"


def main():
    checks = []
    for fmt, content, marker in fixtures():
        evidence = convert_office(
            content,
            format=fmt,
            source_id=f"synthetic:{fmt}",
            source_version="1",
            site_packages=Path(sysconfig.get_paths()["purelib"]),
            expected_version="2.126.0",
        )
        if marker not in json.dumps(evidence["document"]):
            raise ValueError(f"{fmt}: converted content missing")
        if fmt in {"docx", "xlsx"}:
            cells = [
                cell.get("text", "")
                for table in evidence["document"].get("tables", [])
                for cell in table.get("data", {}).get("table_cells", [])
            ]
            if marker not in cells or "128" not in cells:
                raise ValueError(f"{fmt}: structured table cells missing")
            tables = evidence['document'].get('tables', [])
            paired = False
            for table in tables:
                table_cells = table.get('data', {}).get('table_cells', [])
                keys = [cell for cell in table_cells if cell.get('text') == marker]
                values = [cell for cell in table_cells if cell.get('text') == '128']
                paired |= any(key.get('start_row_offset_idx') == value.get('start_row_offset_idx')
                    and key.get('start_row_offset_idx') is not None
                    and value.get('start_col_offset_idx') == key.get('start_col_offset_idx', -2) + 1
                    for key in keys for value in values)
            if not paired:
                raise ValueError(f'{fmt}: parameter/value association lost')
        if evidence["status"] != "unreviewed" or evidence["automatic_reply_eligible"]:
            raise ValueError("synthetic content cannot become reviewed knowledge")
        if not evidence["reading_order"]["body"]:
            raise ValueError(f"{fmt}: missing reading order")
        verify_saved_draft(evidence)
        if fmt == "docx":
            texts = [item.get('text', '') for item in evidence['document'].get('texts', [])]
            if 'printenv bootcmd\nprintenv bootargs\necho "${fixture_value}"' not in texts:
                raise ValueError('docx: command newlines, quotes or variable syntax changed')
            if '合成命令示例（不要执行）' not in texts:
                raise ValueError('docx: Chinese warning heading lost')
            with tempfile.TemporaryDirectory(
                prefix="codex-attachment-canary-"
            ) as directory:
                source = Path(directory) / "synthetic.docx"
                output = Path(directory) / "evidence.json"
                source.write_bytes(content)
                with contextlib.redirect_stdout(io.StringIO()):
                    code = attachment_main(
                        [
                            "--input",
                            str(source),
                            "--output",
                            str(output),
                            "--source-id",
                            "synthetic:cli",
                            "--source-version",
                            "1",
                            "--parser-version",
                            "2.126.0",
                        ]
                    )
                if code != 0 or output.stat().st_mode & 0o777 != 0o600:
                    raise ValueError("private CLI export failed")
                view = preview(json.loads(output.read_text()))
                if (
                    not view["read_only"]
                    or view["status"] != "unreviewed"
                    or not view["items"]
                ):
                    raise ValueError("CLI artifact to review preview failed")
        checks.append(
            {
                "format": fmt,
                "input_bytes": len(content),
                "nodes": len(evidence["reading_order"]["body"]),
                "source_hash_verified": evidence["conversion"]["source_hash_verified"],
                "table_row_association_checked": fmt in {'docx', 'xlsx'},
                "literal_command_block_checked": fmt == 'docx',
                "saved_reopened_exported_unreviewed": True,
            }
        )
    print(
        json.dumps(
            {
                "ok": True,
                "live_data_used": False,
                "cli_export_to_review": True,
                "checks": checks,
            }
        )
    )


if __name__ == "__main__":
    main()
